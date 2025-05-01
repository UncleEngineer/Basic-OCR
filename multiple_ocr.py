from PIL import Image
import pytesseract
import os

# img_doc = Image.open('fbpost.png') # 'C:\\users\\uncleengineer\\desktop\\basic ocr\\eng_text.png'
# ocr_text = pytesseract.image_to_string(img_doc,lang='tha+eng')
# print(ocr_text)
PATH = r'C:\Users\Asus\Desktop\RPA 2025\Basic OCR'

listfiles = os.listdir(PATH)
# print(listfiles)

filter_files = []

for f in listfiles:
    img = os.path.join(PATH,f)
    #print(img)
    if f[:3] == 'bbc':
        filter_files.append(img)

text = ''

for f in filter_files:
    img_doc = Image.open(f)
    ocr_text = pytesseract.image_to_string(img_doc,lang='tha+eng')
    print(ocr_text)
    text += '\n\n' + ocr_text
    print('--------')


from docx import Document
document = Document()
document.add_heading('BBC', 0)

p = document.add_paragraph(text)

document.save('bbc-news2.docx')